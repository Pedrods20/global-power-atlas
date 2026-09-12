"""Official monthly fuel, carbon and FX references for European spreads.

These inputs are deliberately kept outside the canonical interval store: they
are monthly reference series, not electricity-market settlement intervals.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from pathlib import Path

import httpx
import polars as pl
from openpyxl import load_workbook

WORLD_BANK_URL = (
    "https://thedocs.worldbank.org/en/doc/"
    "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/"
    "CMO-Historical-Data-Monthly.xlsx"
)
EEX_CURRENT_URL = (
    "https://public.eex-group.com/eex/eua-auction-report/"
    "emission-spot-primary-market-auction-report-2026-data.xlsx"
)
EEX_HISTORY_URL = (
    "https://www.eex.com/fileadmin/EEX/Downloads/Markets/Environmentals/"
    "EUA_Emission_Spot_Primary_Market_Auction_Report/Archive_Reports/"
    "emission-spot-primary-market-auction-report-2012-2025-data.zip"
)
ECB_URL = (
    "https://data-api.ecb.europa.eu/service/data/EXR/M.USD.EUR.SP00.A"
    "?format=csvdata&startPeriod=2024-01"
)

# Transparent engineering assumptions for historical, indicative spreads.
MMBTU_PER_MWHTH = 3.412141633
COAL_MWHTH_PER_TONNE = 6.978  # 6,000 kcal/kg thermal coal
GAS_EFFICIENCY = 0.50
COAL_EFFICIENCY = 0.38
GAS_TCO2_PER_MWHTH = 0.20196
COAL_TCO2_PER_MWHTH = 0.34056


def reference_path() -> Path:
    return Path(__file__).resolve().parents[2] / "data" / "reference" / "europe_benchmarks.parquet"


def _get(client: httpx.Client, url: str) -> bytes:
    response = client.get(url)
    response.raise_for_status()
    return response.content


def parse_world_bank(content: bytes) -> pl.DataFrame:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook["Monthly Prices"]
    headers = [cell.value for cell in sheet[5]]
    gas_index = headers.index("Natural gas, Europe")
    coal_index = headers.index("Coal, Australian")
    rows = []
    for row in sheet.iter_rows(min_row=7, values_only=True):
        period = row[0]
        if not isinstance(period, str) or not re.fullmatch(r"\d{4}M\d{2}", period):
            continue
        gas, coal = row[gas_index], row[coal_index]
        if isinstance(gas, (int, float)) and isinstance(coal, (int, float)):
            rows.append(
                {
                    "month": f"{period[:4]}-{period[-2:]}",
                    "gas_usd_mmbtu": float(gas),
                    "coal_usd_tonne": float(coal),
                }
            )
    return pl.DataFrame(rows).sort("month")


def parse_eex_workbook(content: bytes) -> pl.DataFrame:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = []
    for row in sheet.iter_rows(min_row=8, values_only=True):
        date, contract, status, price, volume = row[1], row[4], row[5], row[6], row[11]
        if (
            isinstance(date, (dt.date, dt.datetime))
            and isinstance(contract, str)
            and contract.startswith("T")
            and str(status).lower() == "successful"
            and isinstance(price, (int, float))
            and isinstance(volume, (int, float))
        ):
            rows.append(
                {
                    "month": date.strftime("%Y-%m"),
                    "price": float(price),
                    "volume": float(volume),
                }
            )
    return pl.DataFrame(rows)


def parse_eex(current: bytes, history: bytes) -> pl.DataFrame:
    frames = [parse_eex_workbook(current)]
    with zipfile.ZipFile(io.BytesIO(history)) as archive:
        for name in archive.namelist():
            if name.endswith(("2024-data.xlsx", "2025-data.xlsx")):
                frames.append(parse_eex_workbook(archive.read(name)))
    auctions = pl.concat(frames)
    return (
        auctions.group_by("month")
        .agg(
            ((pl.col("price") * pl.col("volume")).sum() / pl.col("volume").sum()).alias(
                "eua_eur_tco2"
            )
        )
        .sort("month")
    )


def parse_ecb(content: bytes) -> pl.DataFrame:
    rows = [
        {"month": row["TIME_PERIOD"], "usd_per_eur": float(row["OBS_VALUE"])}
        for row in csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        if row.get("TIME_PERIOD") and row.get("OBS_VALUE")
    ]
    return pl.DataFrame(rows).sort("month")


def refresh(path: Path | None = None) -> pl.DataFrame:
    """Download official inputs, align them monthly and persist a compact table."""
    with httpx.Client(timeout=60, follow_redirects=True) as client:
        world_bank = parse_world_bank(_get(client, WORLD_BANK_URL))
        eex = parse_eex(_get(client, EEX_CURRENT_URL), _get(client, EEX_HISTORY_URL))
        ecb = parse_ecb(_get(client, ECB_URL))
    frame = (
        world_bank.join(ecb, on="month", how="inner")
        .join(eex, on="month", how="inner")
        .filter(pl.col("month") >= "2024-09")
        .with_columns(
            (pl.col("gas_usd_mmbtu") * MMBTU_PER_MWHTH / pl.col("usd_per_eur")).alias(
                "gas_eur_mwhth"
            ),
            (pl.col("coal_usd_tonne") / pl.col("usd_per_eur")).alias("coal_eur_tonne"),
        )
        .sort("month")
    )
    destination = path or reference_path()
    destination.parent.mkdir(parents=True, exist_ok=True)
    frame.write_parquet(destination, compression="zstd", statistics=True)
    return frame


def calculate_spreads(benchmarks: pl.DataFrame, power: pl.DataFrame) -> pl.DataFrame:
    """Calculate historical clean spark/dark spreads from aligned monthly means."""
    if benchmarks.is_empty() or power.is_empty():
        return pl.DataFrame()
    return (
        power.select(pl.col("period").alias("month"), pl.col("all_hours").alias("power_eur_mwh"))
        .join(benchmarks, on="month", how="inner")
        .with_columns(
            (
                pl.col("power_eur_mwh")
                - pl.col("gas_eur_mwhth") / GAS_EFFICIENCY
                - pl.col("eua_eur_tco2") * GAS_TCO2_PER_MWHTH / GAS_EFFICIENCY
            ).alias("clean_spark_eur_mwh"),
            (
                pl.col("power_eur_mwh")
                - pl.col("coal_eur_tonne") / COAL_MWHTH_PER_TONNE / COAL_EFFICIENCY
                - pl.col("eua_eur_tco2") * COAL_TCO2_PER_MWHTH / COAL_EFFICIENCY
            ).alias("clean_dark_eur_mwh"),
        )
        .sort("month")
    )
